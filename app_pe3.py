import os
import shutil
import tempfile
import streamlit as st
from funciones_asistencia import generar_asistencia
from funciones_op1 import generar_op1
from funciones_op2 import generar_op2
from openpyxl import load_workbook
import io

# ---------------- CONFIGURACIÓN ---------------- #
st.set_page_config(page_title="Sistema PE", layout="wide")

# Variables de sesión
for key in ["asistencia_generada", "op1_generada", "op2_generada"]:
    if key not in st.session_state:
        st.session_state[key] = None

# Plantilla base
PLANTILLA_PATH = os.path.join("plantillas", "PE3 - Reporte.xlsx")


# ---------------- FUNCIONES AUXILIARES ---------------- #
def get_temp_copy():
    """Crea una copia temporal de la plantilla PE3 - Reporte.xlsx."""
    tmp = tempfile.NamedTemporaryFile(delete=False, suffix=".xlsx")
    shutil.copyfile(PLANTILLA_PATH, tmp.name)
    return tmp.name


def clasificar_archivos(lista_archivos):
    """Clasifica los archivos según su nombre."""
    resultado = {
        "asc": None, "nom": None, "acc": None,
        "asc_inst": None, "nom_inst": None, "acc_inst": None,
        "asc_fa": None, "acc_fa": None,
    }

    for file in lista_archivos:
        nombre = file.name.upper().replace(" ", "")
        if "POSTULANTE" in nombre:
            if "ASC" in nombre:
                resultado["asc"] = file
            elif "NOM" in nombre:
                resultado["nom"] = file
            elif "ACC" in nombre:
                resultado["acc"] = file
        elif "INSTRUMENTO" in nombre:
            if "ASC" in nombre:
                resultado["asc_inst"] = file
            elif "NOM" in nombre:
                resultado["nom_inst"] = file
            elif "ACC" in nombre:
                resultado["acc_inst"] = file
        elif "FA" in nombre:
            if "ASC" in nombre:
                resultado["asc_fa"] = file
            elif "ACC" in nombre:
                resultado["acc_fa"] = file
    return resultado


def combinar_reportes(plantilla_path, asistencia, op1, op2):
    """Combina los tres reportes (Asistencia, OP1, OP2) en una sola plantilla Excel."""
    wb_final = load_workbook(plantilla_path)

    wb_a = load_workbook(asistencia)
    wb_op1 = load_workbook(op1)
    wb_op2 = load_workbook(op2)

    # Copiar cada hoja de los reportes al archivo final
    for workbook in [wb_a, wb_op1, wb_op2]:
        for ws_name in workbook.sheetnames:
            ws = workbook[ws_name]
            if ws_name not in wb_final.sheetnames:
                target = wb_final.create_sheet(ws_name)
            else:
                target = wb_final[ws_name]
            for r_idx, row in enumerate(ws.iter_rows(values_only=True), 1):
                for c_idx, value in enumerate(row, 1):
                    target.cell(row=r_idx, column=c_idx, value=value)

    output = io.BytesIO()
    wb_final.save(output)
    output.seek(0)
    return output


# ---------------- ESTILO GENERAL ---------------- #
st.markdown("""
<style>
.block-container {
    padding-top: 2.5rem;    /* 🔼 Aumenta el espacio superior */
    padding-bottom: 0.8rem;
}
h2, h3 {
    color: #1E3A8A;
    font-weight: 700;
}
h2 {
    display: flex;
    align-items: center;
    gap: 0.4rem;
}
.stButton>button {
    border-radius: 6px;
    height: 2.2em;
    font-weight: 600;
    font-size: 0.9em;
    padding: 0 1em;
}
.stButton>button:hover {opacity: 0.95;}
div[data-testid="stHorizontalBlock"] {gap: 0.5rem;}
</style>
""", unsafe_allow_html=True)


# ---------------- CABECERA ---------------- #
st.markdown("<h2 style='text-align:center;'>📊 Sistema de Generación de Reportes PE</h2>", unsafe_allow_html=True)
st.caption("Sube los archivos correspondientes (.xlsx) de Postulantes, Instrumentos y FA. La app los clasificará automáticamente.")


# ---------------- CARGA Y CLASIFICACIÓN ---------------- #
with st.expander("📁 Subir o revisar archivos cargados", expanded=True):
    archivos = st.file_uploader(
        "Selecciona los archivos (.xlsx)", type=["xlsx"], accept_multiple_files=True
    )

if archivos:
    clasificados = clasificar_archivos(archivos)

    with st.expander("📄 Archivos detectados y clasificados automáticamente", expanded=False):
        cols = st.columns(3)
        for i, (k, v) in enumerate(clasificados.items()):
            if v:
                with cols[i % 3]:
                    st.markdown(f"✅ **{k.upper()}**<br><small>{v.name}</small>", unsafe_allow_html=True)
else:
    st.info("Sube los archivos .xlsx correspondientes.", icon="📂")


# ---------------- GENERAR REPORTES ---------------- #
if archivos:
    st.markdown("### ⚙️ Generar")

    col1, col2, col3 = st.columns(3)

    with col1:
        if st.button("🟢 Asistencia", use_container_width=True,
                     disabled=not all([clasificados["asc"], clasificados["nom"], clasificados["acc"]])):
            base = get_temp_copy()
            generar_asistencia(base, clasificados["asc"], clasificados["nom"], clasificados["acc"])
            st.toast("Reporte Asistencia generado ✅", icon="✅")

    with col2:
        if st.button("🟦 OP1", use_container_width=True,
                     disabled=not all([clasificados["asc_inst"], clasificados["nom_inst"], clasificados["asc_fa"]])):
            base = get_temp_copy()
            generar_op1(base, clasificados["asc_fa"], clasificados["asc_inst"], clasificados["nom_inst"])
            st.toast("Reporte OP1 generado ✅", icon="✅")

    with col3:
        if st.button("🟣 OP2", use_container_width=True,
                     disabled=not all([clasificados["acc_inst"], clasificados["acc_fa"]])):
            base = get_temp_copy()
            generar_op2(base, clasificados["acc_fa"], clasificados["acc_inst"])
            st.toast("Reporte OP2 generado ✅", icon="✅")


    # ---------------- DESCARGAS ---------------- #
    st.divider()
    st.markdown("### ⬇️ Descargas")

    cols_dl = st.columns(3)
    with cols_dl[0]:
        if st.session_state.get("asistencia_generada"):
            st.download_button("Descargar Asistencia", st.session_state["asistencia_generada"],
                               file_name="PE - Reporte_ASISTENCIA.xlsx", use_container_width=True)
    with cols_dl[1]:
        if st.session_state.get("op1_generada"):
            st.download_button("Descargar OP1", st.session_state["op1_generada"],
                               file_name="PE - Reporte_OP1.xlsx", use_container_width=True)
    with cols_dl[2]:
        if st.session_state.get("op2_generada"):
            st.download_button("Descargar OP2", st.session_state["op2_generada"],
                               file_name="PE - Reporte_OP2.xlsx", use_container_width=True)


    # ---------------- COMBINAR REPORTES ---------------- #
    st.divider()
    st.markdown("### 📘 Combinar Reportes en una sola plantilla")

    if (
        st.session_state.get("asistencia_generada")
        and st.session_state.get("op1_generada")
        and st.session_state.get("op2_generada")
    ):
        combinado = combinar_reportes(
            PLANTILLA_PATH,
            st.session_state["asistencia_generada"],
            st.session_state["op1_generada"],
            st.session_state["op2_generada"],
        )
        st.download_button(
            "⬇️ Descargar Reporte Final (Asistencia + OP1 + OP2)",
            combinado,
            file_name="PE - Reporte_Final.xlsx",
            use_container_width=True,
        )
    else:
        st.info("Genera los tres reportes (Asistencia, OP1, OP2) antes de combinarlos.", icon="ℹ️")

